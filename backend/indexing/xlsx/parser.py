"""
Excel workbook parser — flattens each sheet's rows into text with
`openpyxl`, then chunks it.
"""

from openpyxl import load_workbook

from ..chunking import chunk_text


def parse(file_path: str) -> list[str]:
    workbook = load_workbook(file_path, data_only=True, read_only=True)

    lines = []
    for sheet in workbook.worksheets:
        lines.append(f"Sheet: {sheet.title}")
        for row in sheet.iter_rows(values_only=True):
            cells = [str(cell) for cell in row if cell is not None]
            if cells:
                lines.append(" | ".join(cells))

    workbook.close()
    full_text = "\n".join(lines)
    return chunk_text(full_text)